from pathlib import Path
from shutil import copy2
from docx import Document

ROOT = Path(r"C:\Users\softw\Documents\Aedon")
TEMPLATE = Path(r"C:\Users\softw\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-system-design\assets\reference.docx")
OUTPUT = ROOT / "docs" / "Aedon-system-design-interface.docx"


def set_paragraph(paragraph, text):
    paragraph.text = text


def set_cell(table, row, column, text):
    cell = table.cell(row, column)
    cell.text = text


def set_table(table, rows):
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            set_cell(table, r, c, value)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    p = doc.paragraphs
    # Preserve the template's hierarchy and replace all editorial placeholders.
    replacements = {
        8: "AEDON",
        9: "Aedon Interface and Content Editing System Design",
        22: "Aedon is a visual website builder. This proposal refactors the project dashboard and editor so a person can understand where to start, add content, return to their projects, and edit a selected section without relying on hidden controls.",
        23: "The design keeps the existing project model, editor state, responsive preview, local storage fallback, Supabase synchronization, and optional device image upload. It defines the user-facing flows, state boundaries, error behavior, and rollout work for the interface changes.",
        28: "The editor contains a canvas, section library, properties panel, responsive preview controls, project actions, undo and redo. These capabilities were present, but several critical actions depended on icon-only controls, hover states, or an empty panel without a next step. Project entry and image upload also needed clearer paths.",
        29: "The result is unnecessary hesitation: a new user may not know how to create a project, add a section, return to the project list, or start editing after opening the editor. The interface must make the first useful action visible while preserving the compact visual language of Aedon.",
        33: "Figure 1. Aedon interface interaction flow from project list to selected section editing.",
        40: "1. The user opens the project dashboard and sees existing projects plus a visible New project entry point.",
        41: "2. Creating a project generates a project record through the active persistence adapter and opens its editor route.",
        42: "3. The editor loads the project, shows the page canvas, and exposes a Projetos breadcrumb to return to the dashboard.",
        43: "4. The left panel explains how to add a section. The user can choose a category, click a component, or drag it to the canvas.",
        44: "5. Selecting a section activates the right properties panel. Until a selection exists, the panel explains exactly what to do.",
        45: "6. Changes update editor state and use the existing persistence flow. Device controls change preview width without changing content.",
        46: "7. When an image is chosen from a device, Aedon uploads it to project storage when Supabase is available; otherwise it keeps a size-limited local data URL.",
        54: "Projects remain addressable by project id. A project is visible only to its owner when Supabase is enabled. The dashboard lists the current user's projects and no longer assumes a single project per owner.",
        55: "A section instance retains its variant, properties, styles, visibility, and ordering in project state. Selecting and editing a section changes only that instance unless the user explicitly applies a supported bulk action.",
        56: "Image fields store a usable URL. Cloud uploads use an owner-scoped path in the project-images bucket. Local fallback images are limited to 1.5 MB to keep local project state reliable.",
        57: "Project creation, rename, duplication, deletion, and editor navigation must surface persistence errors to the user and retain the current view when an operation fails.",
        58: "The interface does not infer publish success. Publishing remains a separate capability until a publishing contract is implemented.",
        59: "UI labels, descriptions, and accessible names are part of the contract for the main controls: projects navigation, device preview, project name, add section, and section selection.",
        63: "The editor uses the existing project state and persistence mechanisms. A save retry must never create duplicate projects from the same user action. When persistence is unavailable, the local adapter provides continuity and the UI makes no claim that cloud synchronization completed.",
        67: "Supabase Auth identifies the project owner. Project queries and mutations are scoped to that owner through the existing policies and client context.",
        68: "Device images are accepted only through the application upload flow. Cloud paths are owner-scoped and storage policies restrict access to the matching authenticated user folder.",
        69: "The editor treats uploaded URLs as user content. Rendering must continue to use the existing safe image and link handling rules.",
        70: "Local fallback data URLs remain in the browser project state. The size limit prevents unexpectedly large browser storage usage.",
        71: "No new third-party analytics, tracking identifiers, or public sharing behavior are introduced by this interface refactor.",
        81: "What exact publish workflow, domain configuration, and publish status should the editor expose?",
        82: "Should replacing an uploaded image delete the previous cloud object immediately or retain it for undo and recovery?",
        83: "When a local project later signs in, should Aedon offer an explicit migration to the cloud project store?",
        84: "Should project cards eventually render saved page thumbnails instead of the lightweight structural preview?",
        87: "Adopt the interface refactor as the default Aedon editing flow. Ship the dashboard and editor clarity changes behind the existing project persistence layer, then measure completion of project creation and first section insertion before expanding the publishing experience.",
    }
    for index, value in replacements.items():
        set_paragraph(p[index], value)

    set_table(doc.tables[0], [["STATUS", "PROPOSED", "OWNER", "Aedon Product", "LAST UPDATED: September 18 2026"]])
    set_table(doc.tables[1], [
        ["AUTHORS", "Aedon Product and Engineering"],
        ["REVIEWERS", "Product and Engineering"],
        ["RELATED DOCS", "Project dashboard and editor"],
        ["SCOPE", "Project entry, editor navigation, section insertion, properties guidance, and image upload."],
    ])
    set_table(doc.tables[2], [
        ["GOALS", "NON-GOALS"],
        ["Make project creation and opening obvious from the dashboard.", "Implement a complete publishing platform."],
        ["Make the next editing action clear when no section is selected.", "Change the content model for every section variant."],
        ["Expose direct, accessible controls for navigation and device preview.", "Add real-time collaboration or comments."],
        ["Support device image upload within existing storage and local limits.", "Introduce product analytics or tracking identifiers."],
    ])
    set_table(doc.tables[3], [
        ["COMPONENT", "RESPONSIBILITY", "PRIMARY STORAGE", "FAILURE BEHAVIOR"],
        ["Project dashboard", "Lists, creates, opens, duplicates, renames, and deletes projects.", "Project adapter", "Shows the persistence error and preserves the current list."],
        ["Editor shell", "Coordinates canvas, panels, device preview, navigation, history, export, and publish entry points.", "In-memory editor state", "Keeps the active project view and exposes usable labels."],
        ["Section library", "Explains and adds section variants by click or drag and drop.", "Section catalog", "Leaves state unchanged if a variant cannot be added."],
        ["Properties panel", "Edits the selected section and explains the empty selection state.", "Project section state", "Does not show misleading controls without a selected section."],
        ["Image upload adapter", "Stores cloud images or local fallback data URLs.", "Supabase Storage or local state", "Returns a clear size, auth, bucket, or network error."],
    ])
    set_table(doc.tables[4], [
        ["FIELD", "TYPE", "REQUIRED", "DESCRIPTION"],
        ["project.id", "string", "Yes", "Stable project identifier used in the editor route."],
        ["project.name", "string", "Yes", "Editable from the project card and editor header."],
        ["project.ownerId", "string", "Cloud only", "Authenticated user identifier used by Supabase policies."],
        ["project.pages", "array", "Yes", "Contains ordered page and section data."],
        ["section.variantId", "string", "Yes", "Determines the chosen component variant."],
        ["section.props and styles", "object", "Yes", "Stores section-specific content and appearance."],
        ["image URL", "string", "No", "Cloud URL or limited local data URL."],
    ])
    set_table(doc.tables[5], [
        ["SCENARIO", "EXPECTED BEHAVIOR", "REASONING"],
        ["Repeated create action", "One user action must resolve to one project creation request.", "Prevents duplicate projects from retry or double-click behavior."],
        ["Project save fails", "Show an error and retain the current project state for retry.", "Avoids silent data loss and makes the failure actionable."],
        ["Upload dependency fails", "Do not set a broken image URL; return the upload error.", "Keeps section data valid and supports a retry."],
        ["Project state changes during editing", "The active editor state remains authoritative until the existing persistence layer confirms an update.", "Prevents the interface from claiming synchronization it has not completed."],
    ])
    set_table(doc.tables[6], [
        ["SIGNAL", "SLO OR ALERT", "OWNER", "LAUNCH GATE"],
        ["Project creation", "New project opens a unique project route.", "Engineering", "Required"],
        ["Section insertion", "Click and drag add exactly one intended section.", "Product and Engineering", "Required"],
        ["Image upload", "Accepted image yields a URL within configured size limits.", "Engineering", "Required"],
        ["Navigation", "Projetos returns to the dashboard at every supported width.", "Product", "Required"],
        ["Accessibility", "Primary controls have text or accessible names.", "Engineering", "Required"],
        ["Rollout constraint", "Validate desktop, tablet, mobile, keyboard, and drag interactions before release.", "Product and Engineering", "Required"],
    ])
    set_table(doc.tables[7], [
        ["ALTERNATIVE", "WHY IT WAS CONSIDERED", "WHY IT WAS NOT SELECTED"],
        ["Keep hover-only actions", "Uses less visible space.", "Hides core project actions and harms touch use."],
        ["Use a single floating add button", "Minimal chrome.", "Does not explain available section types or drag behavior."],
        ["Force cloud upload only", "One storage path.", "Removes the existing local development and offline fallback."],
        ["Build publishing with this refactor", "One larger release.", "Publishing needs its own product and technical contract."],
    ])
    set_table(doc.tables[8], [
        ["MILESTONE", "DELIVERABLE", "EXIT CRITERIA"],
        ["M1 Interface clarity", "Dashboard cards, visible actions, navigation breadcrumb, library guidance, and empty properties state.", "A new user can create, open, add, and select without hidden controls."],
        ["M2 Upload validation", "Device image upload backed by Supabase Storage with local fallback.", "Cloud bucket and policies are deployed; errors are understandable."],
        ["M3 Interaction QA", "Desktop, tablet, mobile, keyboard, and drag-and-drop verification.", "No text selection during intended drag interactions."],
        ["M4 Publish design", "Separate publish requirements and API contract.", "Decision made on domains, deployment, and status feedback."],
    ])
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
